"""投递模块测试：服务层 TC-01~04 + API 集成 TC-26 / TC-46（接口文档 3.3）。"""

from datetime import date, datetime, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.exceptions import BizException, ErrorCode
from app.models import Application, InterviewSession, JdAnalysisReport
from app.models.enums import ApplicationStatus, CloseReason
from app.schemas.application import ApplicationCreate, ApplicationStatusUpdate
from app.services import application_service

API = "/api/v1/applications"


def _xlsx(rows: list[list]) -> bytes:
    """把若干行（首行为表头）组装成 xlsx 字节流，供导入用例使用。"""
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------- TC-01 新增投递 ----------


def test_create_application_defaults(db_session: Session, account_id: int):
    """TC-01：字段合法入库，status 默认 APPLIED，投递日期默认当天。"""
    dto = application_service.create_application(
        db_session,
        account_id, ApplicationCreate(company="浩鲸科技", position="Java 开发", city="南京", expected_salary="13k", channel="官网"),
    )

    assert dto.id > 0
    assert dto.status == ApplicationStatus.APPLIED
    assert dto.applied_at == date.today()
    assert (dto.company, dto.position, dto.city) == ("浩鲸科技", "Java 开发", "南京")

    row = db_session.get(Application, dto.id)
    assert row is not None and row.status == "APPLIED"


# ---------- TC-02 状态流转 ----------


def test_change_status_legal_chain(db_session: Session, account_id: int):
    """TC-02：合法流转链 APPLIED→WRITTEN→INTERVIEW→OFFER 逐跳成功。"""
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="浩鲸科技", position="Java 开发")
    )

    for target in (ApplicationStatus.WRITTEN, ApplicationStatus.INTERVIEW, ApplicationStatus.OFFER):
        dto = application_service.change_status(db_session, account_id, dto.id, ApplicationStatusUpdate(status=target))
        assert dto.status == target


@pytest.mark.parametrize(
    ("start", "target"),
    [
        (ApplicationStatus.APPLIED, ApplicationStatus.INTERVIEW),  # 不设笔试的公司
        (ApplicationStatus.APPLIED, ApplicationStatus.OFFER),  # 直通 offer
        (ApplicationStatus.WRITTEN, ApplicationStatus.OFFER),  # 面试环节合并
    ],
)
def test_change_status_legal_skip_level(db_session: Session, account_id: int, start: ApplicationStatus, target: ApplicationStatus):
    """TC-02：跳级前进合法（SRS v1.8 放开）——可前进到链上任意更靠后的状态。"""
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="云器科技", position="后端开发", status=start)
    )

    dto = application_service.change_status(db_session, account_id, dto.id, ApplicationStatusUpdate(status=target))

    assert dto.status == target


@pytest.mark.parametrize(
    "start",
    [
        ApplicationStatus.APPLIED,
        ApplicationStatus.WRITTEN,
        ApplicationStatus.INTERVIEW,
        ApplicationStatus.OFFER,  # 拒 offer / 被撤回 / 谈崩，同样可结束
    ],
)
def test_change_status_closed_from_any_active_state(db_session: Session, account_id: int, start: ApplicationStatus):
    """TC-02：除已结束外任意状态（含 OFFER）均可流转至 CLOSED，结束原因同时入库。"""
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="亚信科技", position="后端开发", status=start)
    )

    dto = application_service.change_status(
        db_session,
        account_id, dto.id,
        ApplicationStatusUpdate(status=ApplicationStatus.CLOSED, close_reason=CloseReason.FAILED),
    )

    assert dto.status == ApplicationStatus.CLOSED
    assert dto.close_reason == CloseReason.FAILED


@pytest.mark.parametrize(
    ("start", "target"),
    [
        (ApplicationStatus.WRITTEN, ApplicationStatus.APPLIED),  # 回退
        (ApplicationStatus.INTERVIEW, ApplicationStatus.WRITTEN),  # 回退
        (ApplicationStatus.OFFER, ApplicationStatus.INTERVIEW),  # 回退
        (ApplicationStatus.APPLIED, ApplicationStatus.APPLIED),  # 原地不动
        (ApplicationStatus.INTERVIEW, ApplicationStatus.INTERVIEW),  # 原地不动
        (ApplicationStatus.CLOSED, ApplicationStatus.WRITTEN),  # 终态不可再流转
        (ApplicationStatus.CLOSED, ApplicationStatus.APPLIED),  # 终态不可再流转
    ],
)
def test_change_status_illegal_rejected(db_session: Session, account_id: int, start: ApplicationStatus, target: ApplicationStatus):
    """TC-02：非法流转拒绝（错误码 10001）——回退 / 原地 / 终态；且库中状态不变。"""
    dto = application_service.create_application(
        db_session,
        account_id, ApplicationCreate(company="新华三", position="Java 开发", status=start),
    )

    with pytest.raises(BizException) as exc_info:
        application_service.change_status(db_session, account_id, dto.id, ApplicationStatusUpdate(status=target))

    assert exc_info.value.code == ErrorCode.PARAM_INVALID
    db_session.expire_all()
    assert db_session.get(Application, dto.id).status == start


def test_change_status_event_and_remark_appended(db_session: Session, account_id: int):
    """TC-02：转 WRITTEN/INTERVIEW 时更新下次考试时间，流转备注追加不覆盖原备注。"""
    event_at = f"{date.today() + timedelta(days=3)} 14:00:00"
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="满帮", position="Java 开发", remark="原备注")
    )

    dto = application_service.change_status(
        db_session,
        account_id, dto.id,
        ApplicationStatusUpdate(status=ApplicationStatus.WRITTEN, event_at=event_at, remark="收到笔试通知"),
    )

    assert dto.next_event_at.strftime("%Y-%m-%d %H:%M:%S") == event_at
    assert dto.remark.startswith("原备注\n[")
    assert "流转至 待笔试] 收到笔试通知" in dto.remark


# ---------- TC-03 批量导入（行级容错） ----------


def test_import_xlsx_partial_rows(db_session: Session, account_id: int):
    """TC-03：xlsx 合法行入库，非法行只记错误不中断其余行。"""
    content = _xlsx(
        [
            ["company", "position", "city", "applied_at", "channel", "status", "remark"],
            ["浩鲸科技", "Java 开发", "南京", "2026-09-01", "官网", "APPLIED", "已投"],
            ["", "前端开发", "南京", "2026-09-02", "BOSS", "APPLIED", ""],  # 公司为空
            ["亚信科技", "后端开发", "上海", "2026-09-03", "内推", "WRITTEN", ""],
        ]
    )

    result = application_service.import_from_file(db_session, account_id, "投递清单.xlsx", content)

    assert result.success_count == 2
    assert [(item.row, item.reason) for item in result.errors] == [(3, "公司名称不能为空")]
    imported = {row.company: row.status for row in db_session.scalars(select(Application)).all()}
    assert imported == {"浩鲸科技": "APPLIED", "亚信科技": "WRITTEN"}


def test_import_csv_partial_rows(db_session: Session, account_id: int):
    """TC-03：csv（GBK 编码）解析成功；非法状态与非法日期进 errors，空白行跳过。"""
    content = (
        "company,position,applied_at,status\n"
        "云器科技,Java 开发,2026/09/01,\n"  # 日期用 / 分隔可解析，状态空走默认
        "合合信息,,2026-09-02,APPLIED\n"  # 岗位为空
        ",,,\n"  # 整行空白：跳过
        "哈啰出行,后端开发,9月3日,APPLIED\n"  # 日期格式非法
        "招银网络,后端开发,2026-09-04,UNKNOWN\n"  # 状态非法
    ).encode("gbk")

    result = application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert result.success_count == 1
    assert [(item.row, item.reason) for item in result.errors] == [
        (3, "岗位名称不能为空"),
        (5, "投递日期格式应为 YYYY-MM-DD"),
        (6, "状态值非法（可选：已投递、待笔试、面试中、已获 offer、已结束，或对应英文枚举）"),
    ]


def test_import_all_rows_invalid_raises_20002(db_session: Session, account_id: int):
    """TC-03：全部行非法时不入库，返回 20002 并携带行级清单。"""
    content = "company,position\n,Java\n,后端\n".encode("utf-8")

    with pytest.raises(BizException) as exc_info:
        application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert exc_info.value.code == ErrorCode.IMPORT_ROW_INVALID
    assert exc_info.value.data["success_count"] == 0
    assert [item["row"] for item in exc_info.value.data["errors"]] == [2, 3]
    assert db_session.query(Application).count() == 0


# ---------- TC-04 文件与表头校验 ----------


def test_import_missing_required_header(db_session: Session, account_id: int):
    """TC-04：表头缺少必填列（岗位）→ 20001。"""
    with pytest.raises(BizException) as exc_info:
        application_service.import_from_file(db_session, account_id, "import.csv", b"company,city\nA,\xe5\x8d\x97\xe4\xba\xac\n")

    assert exc_info.value.code == ErrorCode.IMPORT_FILE_INVALID
    assert "岗位" in exc_info.value.message


def test_import_empty_header(db_session: Session, account_id: int):
    """TC-04：没有表头行（空文件）→ 20001。"""
    with pytest.raises(BizException) as exc_info:
        application_service.import_from_file(db_session, account_id, "import.csv", b"\n")

    assert exc_info.value.code == ErrorCode.IMPORT_FILE_INVALID


def test_import_unsupported_file_type(db_session: Session, account_id: int):
    """TC-04：扩展名非 xlsx/csv → 20001。"""
    with pytest.raises(BizException) as exc_info:
        application_service.import_from_file(db_session, account_id, "import.pdf", b"whatever")

    assert exc_info.value.code == ErrorCode.IMPORT_FILE_INVALID


# ---------- 删除的引用维护 ----------


def test_delete_application_detaches_related_records(db_session: Session, account_id: int):
    """删除投递：关联 JD 报告/面试会话保留数据、解除关联（否则外键拦截删除会 500）。"""
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="云器科技", position="后端开发")
    )
    report = JdAnalysisReport(user_id=account_id, application_id=dto.id, jd_text="JD 原文", report_text="五段报告")
    session = InterviewSession(user_id=account_id, application_id=dto.id, company="云器科技", position="后端开发")
    db_session.add_all([report, session])
    db_session.commit()

    application_service.delete_application(db_session, account_id, dto.id)

    db_session.expire_all()
    assert db_session.get(Application, dto.id) is None
    assert db_session.get(JdAnalysisReport, report.id).application_id is None
    assert db_session.get(InterviewSession, session.id).application_id is None


# ---------- TC-26 投递 CRUD 全流程（API 集成） ----------


def test_application_crud_flow(client: TestClient):
    """TC-26：创建→详情→编辑→列表筛选分页→删除→404，全链路走接口文档口径。"""
    created = client.post(
        API,
        json={
            "company": "浩鲸科技",
            "position": "Java 开发",
            "city": "南京",
            "expected_salary": "13k",
            "applied_at": "2026-09-20",
            "channel": "官网",
            "remark": "面试已过 BP 面",
        },
    )
    assert created.status_code == 200
    body = created.json()
    assert body["code"] == 0 and body["message"] == "ok"
    app_id = body["data"]["id"]
    assert body["data"]["status"] == "APPLIED"
    assert body["data"]["applied_at"] == "2026-09-20"
    assert body["data"]["created_at"].count(":") == 2 and "T" not in body["data"]["created_at"]

    # 详情：含备注全文
    detail = client.get(f"{API}/{app_id}").json()["data"]
    assert detail["remark"] == "面试已过 BP 面"

    # 编辑（全量更新：未传字段置空，status 不受影响）
    updated = client.put(
        f"{API}/{app_id}",
        json={"company": "浩鲸科技国际", "position": "Java 开发", "status": "OFFER"},
    ).json()["data"]
    assert updated["company"] == "浩鲸科技国际"
    assert updated["city"] is None and updated["expected_salary"] is None
    assert updated["applied_at"] == date.today().isoformat()  # 未传投递日期 → 默认当天
    assert updated["status"] == "APPLIED"  # status 只能走流转接口

    # 状态流转
    patched = client.patch(
        f"{API}/{app_id}/status", json={"status": "WRITTEN", "event_at": "2026-09-26 14:00:00", "remark": "笔试通知"}
    ).json()["data"]
    assert patched["status"] == "WRITTEN"
    assert patched["next_event_at"] == "2026-09-26 14:00:00"

    # 列表：筛选 + 分页 + 列表项不含 remark
    listed = client.get(API, params={"status": "WRITTEN", "company": "浩鲸", "page": 1, "page_size": 10}).json()
    assert listed["data"]["total"] == 1
    assert "remark" not in listed["data"]["items"][0]
    assert client.get(API, params={"status": "OFFER"}).json()["data"]["total"] == 0
    assert client.get(API, params={"page_size": 51}).json()["code"] == 10001  # 分页上限 50

    # 删除 → 详情 404
    assert client.delete(f"{API}/{app_id}").json() == {"code": 0, "message": "ok", "data": None}
    assert client.get(f"{API}/{app_id}").json()["code"] == 10002


def test_application_invalid_payload(client: TestClient):
    """TC-26：必填字段缺失/超长 → 参数校验失败 10001。"""
    assert client.post(API, json={"position": "Java 开发"}).json()["code"] == 10001
    assert client.post(API, json={"company": "  ", "position": "Java 开发"}).json()["code"] == 10001
    assert client.post(API, json={"company": "A" * 101, "position": "Java"}).json()["code"] == 10001
    assert client.get(f"{API}/9999999").json()["code"] == 10002


def test_close_reason_api_flow(client: TestClient):
    """TC-26：结束原因接口口径——缺 reason 结束被拒 10001、带 reason 成功且可回读、列表按 reason 筛选。"""
    app_id = client.post(API, json={"company": "亚信安全", "position": "Java 开发"}).json()["data"]["id"]

    assert client.patch(f"{API}/{app_id}/status", json={"status": "CLOSED"}).json()["code"] == 10001

    closed = client.patch(
        f"{API}/{app_id}/status", json={"status": "CLOSED", "close_reason": "DECLINED"}
    ).json()["data"]
    assert closed["status"] == "CLOSED"
    assert closed["close_reason"] == "DECLINED"

    assert client.patch(f"{API}/{app_id}/status", json={"status": "INTERVIEW"}).json()["code"] == 10001  # 终态不可再流转
    assert client.get(API, params={"close_reason": "FAILD"}).json()["code"] == 10001  # 非法枚举值
    assert client.get(API, params={"close_reason": "DECLINED"}).json()["data"]["total"] == 1
    assert client.get(f"{API}/{app_id}").json()["data"]["close_reason"] == "DECLINED"


def test_status_skip_level_api(client: TestClient):
    """TC-02：跳级流转走 HTTP 链路（前端看板拖拽场景）——APPLIED 直跳 INTERVIEW 成功、回退仍被拒。"""
    app_id = client.post(API, json={"company": "云器科技", "position": "后端开发"}).json()["data"]["id"]

    jumped = client.patch(f"{API}/{app_id}/status", json={"status": "INTERVIEW"}).json()["data"]
    assert jumped["status"] == "INTERVIEW"

    assert client.patch(f"{API}/{app_id}/status", json={"status": "WRITTEN"}).json()["code"] == 10001  # 回退


def test_import_endpoints_api(client: TestClient):
    """TC-26：模板下载返回 xlsx 文件流；上传导入成功与全非法（20002 带行级清单）。"""
    template = client.get(f"{API}/template")
    assert template.status_code == 200
    assert template.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "application_import_template.xlsx" in template.headers["content-disposition"]
    assert template.content[:2] == b"PK"  # xlsx 即 zip 包

    ok = client.post(
        f"{API}/import",
        files={"file": ("投递清单.xlsx", _xlsx([["company", "position"], ["满帮", "Java 开发"]]), "application/vnd.ms-excel")},
    ).json()
    assert ok["code"] == 0
    assert ok["data"] == {"success_count": 1, "errors": []}

    bad = client.post(f"{API}/import", files={"file": ("bad.csv", b"company,position\n,Java\n", "text/csv")})
    assert bad.status_code == 400
    assert bad.json()["code"] == 20002
    assert bad.json()["data"]["errors"][0] == {"row": 2, "reason": "公司名称不能为空"}


# ---------- TC-46 投递趋势 ----------


def test_trend_zero_filled_and_bounds(client: TestClient):
    """TC-46：按日 GROUP BY 计数、补齐 0 值日期、days 边界（默认 30 / 上限 90）。"""
    today = date.today()
    yesterday = today - timedelta(days=1)
    for company, applied in (("A 公司", today), ("B 公司", today), ("C 公司", yesterday)):
        client.post(API, json={"company": company, "position": "Java 开发", "applied_at": applied.isoformat()})

    items = client.get(f"{API}/trend", params={"days": 7}).json()["data"]["items"]
    assert len(items) == 7
    assert items[0]["date"] == (today - timedelta(days=6)).isoformat()
    assert items[-1] == {"date": today.isoformat(), "count": 2}
    assert items[-2]["count"] == 1
    assert items[0]["count"] == 0  # 无投递日期补 0

    assert len(client.get(f"{API}/trend").json()["data"]["items"]) == 30  # 默认 30 天
    assert client.get(f"{API}/trend", params={"days": 90}).json()["code"] == 0
    assert client.get(f"{API}/trend", params={"days": 91}).json()["code"] == 10001
    assert client.get(f"{API}/trend", params={"days": 0}).json()["code"] == 10001


# ---------- TC-01~04 补强：实现分支覆盖 ----------
#
# 现有用例已覆盖各 TC 的验收点直系路径；本段落补齐 application_service.py 中
# 未被走到的容错与边界分支。
# 分组：A=TC-01 / B=TC-02 / C=TC-04 / D=TC-03 / E=TC-03·TC-04。


# ---------- A 组：新增投递 ----------


def test_create_application_all_fields(db_session: Session, account_id: int):
    """TC-01：全字段显式传入后读回逐字段一致，显式投递日期不被当天覆盖。"""
    dto = application_service.create_application(
        db_session,
        account_id, ApplicationCreate(
            company="亚信科技",
            position="后端开发",
            city="上海",
            expected_salary="15k*14",
            applied_at=date(2026, 9, 1),
            channel="内推",
            status=ApplicationStatus.INTERVIEW,
            next_event_at=datetime(2026, 9, 26, 14, 0, 0),
            remark="二面待约",
        ),
    )

    assert (dto.company, dto.position, dto.city) == ("亚信科技", "后端开发", "上海")
    assert dto.expected_salary == "15k*14"
    assert dto.applied_at == date(2026, 9, 1)
    assert dto.channel == "内推"
    assert dto.status == ApplicationStatus.INTERVIEW
    assert dto.next_event_at == datetime(2026, 9, 26, 14, 0, 0)
    assert dto.remark == "二面待约"

    row = db_session.get(Application, dto.id)
    assert row.applied_at == datetime(2026, 9, 1, 0, 0, 0)  # 库内为 datetime，取当天 00:00:00
    assert row.next_event_at == datetime(2026, 9, 26, 14, 0, 0)


# ---------- B 组：状态流转 ----------


def test_change_status_event_at_ignored_for_closed(db_session: Session, account_id: int):
    """TC-02：event_at 仅在转为 WRITTEN/INTERVIEW 时生效——转 CLOSED 时 next_event_at 保持原值。"""
    original = datetime(2026, 9, 26, 14, 0, 0)
    dto = application_service.create_application(
        db_session,
        account_id, ApplicationCreate(company="满帮", position="Java 开发", next_event_at=original),
    )

    dto = application_service.change_status(
        db_session,
        account_id, dto.id,
        ApplicationStatusUpdate(
            status=ApplicationStatus.CLOSED,
            close_reason=CloseReason.EXPIRED,
            event_at=datetime(2026, 10, 1, 9, 0, 0),
        ),
    )

    assert dto.status == ApplicationStatus.CLOSED
    assert dto.close_reason == CloseReason.EXPIRED
    assert dto.next_event_at == original


def test_change_status_event_at_applied_for_interview(db_session: Session, account_id: int):
    """TC-02：转 INTERVIEW 时 event_at 同样生效（现有用例只覆盖了 WRITTEN）。"""
    dto = application_service.create_application(
        db_session,
        account_id, ApplicationCreate(company="新华三", position="Java 开发", status=ApplicationStatus.WRITTEN),
    )
    event_at = datetime(2026, 9, 28, 10, 30, 0)

    dto = application_service.change_status(
        db_session,
        account_id, dto.id,
        ApplicationStatusUpdate(status=ApplicationStatus.INTERVIEW, event_at=event_at),
    )

    assert dto.status == ApplicationStatus.INTERVIEW
    assert dto.next_event_at == event_at


def test_change_status_remark_without_existing(db_session: Session, account_id: int):
    """TC-02：原备注为空时流转备注直接落库，不产生前导换行。"""
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="云器科技", position="后端开发")
    )

    dto = application_service.change_status(
        db_session,
        account_id, dto.id,
        ApplicationStatusUpdate(status=ApplicationStatus.WRITTEN, remark="收到笔试通知"),
    )

    assert dto.remark is not None
    assert not dto.remark.startswith("\n")
    assert dto.remark.startswith("[")
    assert dto.remark.endswith("流转至 待笔试] 收到笔试通知")


def test_change_status_without_remark_keeps_remark(db_session: Session, account_id: int):
    """TC-02：流转不传 remark 时备注原样保留，不追加任何流转记录。"""
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="合合信息", position="后端开发", remark="原备注")
    )

    dto = application_service.change_status(
        db_session, account_id, dto.id, ApplicationStatusUpdate(status=ApplicationStatus.WRITTEN)
    )

    assert dto.remark == "原备注"


def test_change_status_not_found(db_session: Session, account_id: int):
    """TC-02：对不存在的投递流转 → 10002，且不产生任何写入。"""
    with pytest.raises(BizException) as exc_info:
        application_service.change_status(
            db_session, account_id, 999_999, ApplicationStatusUpdate(status=ApplicationStatus.WRITTEN)
        )

    assert exc_info.value.code == ErrorCode.NOT_FOUND
    assert db_session.query(Application).count() == 0


# ---------- B 组补充：结束原因 close_reason（除已结束外任意状态可结束） ----------


def test_change_status_closed_requires_reason(db_session: Session, account_id: int):
    """TC-02：流转至 CLOSED 未指明结束原因 → 10001，且库中状态不变。"""
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="神州信息", position="Java 开发")
    )

    with pytest.raises(BizException) as exc_info:
        application_service.change_status(
            db_session, account_id, dto.id, ApplicationStatusUpdate(status=ApplicationStatus.CLOSED)
        )

    assert exc_info.value.code == ErrorCode.PARAM_INVALID
    db_session.expire_all()
    row = db_session.get(Application, dto.id)
    assert row.status == ApplicationStatus.APPLIED
    assert row.close_reason is None


def test_change_status_closed_writes_reason(db_session: Session, account_id: int):
    """TC-02：流转至 CLOSED 时结束原因入库，DTO 与库内一致。"""
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="东软集团", position="Java 开发")
    )

    dto = application_service.change_status(
        db_session,
        account_id, dto.id,
        ApplicationStatusUpdate(status=ApplicationStatus.CLOSED, close_reason=CloseReason.DECLINED),
    )

    assert dto.status == ApplicationStatus.CLOSED
    assert dto.close_reason == CloseReason.DECLINED
    assert db_session.get(Application, dto.id).close_reason == "DECLINED"


def test_change_status_active_has_no_reason(db_session: Session, account_id: int):
    """TC-02：非 CLOSED 态的 close_reason 恒为 null（合法链路上逐跳核对）。"""
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="满帮集团", position="Java 开发")
    )

    for target in (ApplicationStatus.WRITTEN, ApplicationStatus.INTERVIEW, ApplicationStatus.OFFER):
        dto = application_service.change_status(db_session, account_id, dto.id, ApplicationStatusUpdate(status=target))
        assert dto.close_reason is None


def test_create_close_reason_kept_only_when_closed(db_session: Session, account_id: int):
    """TC-01：新增时结束原因仅 CLOSED 态入库，非 CLOSED 态传该字段被忽略（不报错）。"""
    closed = application_service.create_application(
        db_session,
        account_id, ApplicationCreate(
            company="汉得信息",
            position="Java 开发",
            status=ApplicationStatus.CLOSED,
            close_reason=CloseReason.EXPIRED,
        ),
    )
    active = application_service.create_application(
        db_session,
        account_id, ApplicationCreate(
            company="华苏科技",
            position="Java 开发",
            status=ApplicationStatus.APPLIED,
            close_reason=CloseReason.FAILED,
        ),
    )

    assert closed.close_reason == CloseReason.EXPIRED
    assert active.close_reason is None
    assert db_session.get(Application, active.id).close_reason is None


def test_update_close_reason_rejected_when_not_closed(db_session: Session, account_id: int):
    """TC-26：编辑接口对非 CLOSED 记录传结束原因 → 10001，且不产生部分更新。"""
    dto = application_service.create_application(
        db_session, account_id, ApplicationCreate(company="焦点科技", position="Java 开发", city="南京")
    )

    with pytest.raises(BizException) as exc_info:
        application_service.update_application(
            db_session,
            account_id, dto.id,
            ApplicationCreate(company="焦点科技", position="Java 开发", close_reason=CloseReason.FAILED),
        )

    assert exc_info.value.code == ErrorCode.PARAM_INVALID
    db_session.expire_all()
    row = db_session.get(Application, dto.id)
    assert row.city == "南京" and row.close_reason is None


def test_update_close_reason_allowed_when_closed(db_session: Session, account_id: int):
    """TC-26：已结束记录可更正结束原因（原因选错时的修正路径）。"""
    dto = application_service.create_application(
        db_session,
        account_id, ApplicationCreate(
            company="润和软件",
            position="Java 开发",
            status=ApplicationStatus.CLOSED,
            close_reason=CloseReason.FAILED,
        ),
    )

    dto = application_service.update_application(
        db_session,
        account_id, dto.id,
        ApplicationCreate(company="润和软件", position="Java 开发", close_reason=CloseReason.DECLINED),
    )

    assert dto.status == ApplicationStatus.CLOSED
    assert dto.close_reason == CloseReason.DECLINED


def test_list_filter_by_close_reason(db_session: Session, account_id: int):
    """TC-26：列表按结束原因过滤，命中项 close_reason 一致，未结束记录不计入。"""
    expected = {
        "南京银行": CloseReason.FAILED,
        "江苏电信": CloseReason.DECLINED,
        "莱斯信息": CloseReason.EXPIRED,
    }
    for company, reason in expected.items():
        application_service.create_application(
            db_session,
            account_id, ApplicationCreate(
                company=company,
                position="Java 开发",
                status=ApplicationStatus.CLOSED,
                close_reason=reason,
            ),
        )
    application_service.create_application(
        db_session, account_id, ApplicationCreate(company="富士通南大", position="Java 开发")
    )

    for company, reason in expected.items():
        page = application_service.list_applications(db_session, user_id=account_id, close_reason=reason)
        assert page.total == 1
        assert page.items[0].company == company
        assert page.items[0].close_reason == reason

    assert application_service.list_applications(db_session, user_id=account_id).total == 4


# ---------- C 组：导入文件级校验（20001 的其余分支） ----------


def test_import_empty_content(db_session: Session, account_id: int):
    """TC-04：文件内容为空（0 字节）→ 20001。"""
    with pytest.raises(BizException) as exc_info:
        application_service.import_from_file(db_session, account_id, "import.csv", b"")

    assert exc_info.value.code == ErrorCode.IMPORT_FILE_INVALID
    assert "文件内容为空" in exc_info.value.message


def test_import_oversized_file(db_session: Session, account_id: int):
    """TC-04：超过 2MB 上限 → 20001（表头合法，仅体积越界）。"""
    content = b"company,position\n" + b"a,b\n" * (2 * 1024 * 1024 // 4 + 1)

    with pytest.raises(BizException) as exc_info:
        application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert exc_info.value.code == ErrorCode.IMPORT_FILE_INVALID
    assert "2MB" in exc_info.value.message
    assert db_session.query(Application).count() == 0


def test_import_uppercase_suffix(db_session: Session, account_id: int):
    """TC-04：扩展名大小写不敏感——.XLSX 与 .xlsx 同样可导入。"""
    content = _xlsx([["company", "position"], ["招银网络", "后端开发"]])

    result = application_service.import_from_file(db_session, account_id, "投递清单.XLSX", content)

    assert result.success_count == 1
    assert result.errors == []
    assert db_session.scalar(select(Application).where(Application.company == "招银网络")) is not None


def test_import_broken_xlsx(db_session: Session, account_id: int):
    """TC-04：.xlsx 扩展名但内容不是合法 xlsx → 20001（而非未捕获异常导致 500）。"""
    with pytest.raises(BizException) as exc_info:
        application_service.import_from_file(db_session, account_id, "broken.xlsx", b"this is not a zip file")

    assert exc_info.value.code == ErrorCode.IMPORT_FILE_INVALID
    assert "无法解析" in exc_info.value.message


def test_import_empty_xlsx(db_session: Session, account_id: int):
    """TC-04：合法 xlsx 但无任何行（连表头都没有）→ 20001。"""
    workbook = Workbook()
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(BizException) as exc_info:
        application_service.import_from_file(db_session, account_id, "empty.xlsx", buffer.getvalue())

    assert exc_info.value.code == ErrorCode.IMPORT_FILE_INVALID
    assert "文件为空" in exc_info.value.message


def test_import_undecodable_csv(db_session: Session, account_id: int):
    """TC-04：既非 UTF-8 也非 GBK 的字节序列 → 20001（不抛 UnicodeDecodeError）。"""
    with pytest.raises(BizException) as exc_info:
        application_service.import_from_file(db_session, account_id, "import.csv", b"\xff\xfeA")

    assert exc_info.value.code == ErrorCode.IMPORT_FILE_INVALID
    assert "编码" in exc_info.value.message


# ---------- D 组：导入解析容错 ----------


def test_import_xlsx_date_cell(db_session: Session, account_id: int):
    """TC-03：xlsx 的 applied_at 为真日期单元格（非文本）时，归一化为 YYYY-MM-DD 后入库。

    实测 openpyxl 读回日期单元格得到 datetime 对象，经 _cell_text 归一化后走文本解析路径。
    """
    content = _xlsx([["company", "position", "applied_at"], ["哈啰出行", "后端开发", date(2026, 9, 1)]])

    result = application_service.import_from_file(db_session, account_id, "投递清单.xlsx", content)

    assert result.success_count == 1
    assert result.errors == []
    row = db_session.scalar(select(Application).where(Application.company == "哈啰出行"))
    assert row.applied_at == datetime(2026, 9, 1, 0, 0, 0)


def test_import_csv_with_bom(db_session: Session, account_id: int):
    """TC-03：带 UTF-8 BOM 的 csv（Excel 另存常见）表头不被 BOM 污染，解析成功。"""
    content = "company,position\n满帮,Java 开发\n".encode("utf-8-sig")

    result = application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert result.success_count == 1
    assert result.errors == []


def test_import_extra_and_missing_columns(db_session: Session, account_id: int):
    """TC-03：表头含模板外多余列时忽略该列；行内缺列按空串处理（缺必填列则该行报错）。"""
    content = (
        "company,position,city,status,note\n"
        "云器科技,后端开发,南京,APPLIED,多余列内容应被忽略\n"
        "合合信息\n"
    ).encode("utf-8")

    result = application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert result.success_count == 1
    assert [(item.row, item.reason) for item in result.errors] == [(3, "岗位名称不能为空")]
    row = db_session.scalar(select(Application).where(Application.company == "云器科技"))
    assert (row.city, row.status) == ("南京", "APPLIED")


def test_import_applied_at_with_time(db_session: Session, account_id: int):
    """TC-03：applied_at 带时间部分时取日期，不判为格式非法。"""
    content = "company,position,applied_at\n满帮,Java 开发,2026-09-01 10:30:00\n".encode("utf-8")

    result = application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert result.success_count == 1
    assert result.errors == []
    row = db_session.scalar(select(Application).where(Application.company == "满帮"))
    assert row.applied_at == datetime(2026, 9, 1, 0, 0, 0)


# ---------- E 组：表头与行级校验 ----------


def test_import_missing_multiple_required_headers(db_session: Session, account_id: int):
    """TC-04：表头同时缺「公司」与「岗位」时，提示一次报全两列名。"""
    with pytest.raises(BizException) as exc_info:
        application_service.import_from_file(db_session, account_id, "import.csv", "city,channel\n南京,官网\n".encode("utf-8"))

    assert exc_info.value.code == ErrorCode.IMPORT_FILE_INVALID
    assert "公司" in exc_info.value.message
    assert "岗位" in exc_info.value.message


@pytest.mark.parametrize(
    ("column", "value", "reason"),
    [
        ("company", "A" * 101, "公司名称超过 100 字"),
        ("position", "B" * 101, "岗位名称超过 100 字"),
        ("city", "C" * 51, "城市超过 50 字"),
        ("channel", "D" * 51, "投递渠道超过 50 字"),
    ],
)
def test_import_row_field_too_long(db_session: Session, account_id: int, column: str, value: str, reason: str):
    """TC-03：行级字段超长只废该行，其余合法行照常入库。"""
    invalid = {
        "company": "满帮",
        "position": "Java 开发",
        "city": "南京",
        "applied_at": "",
        "channel": "官网",
        "status": "",
        "remark": "",
    }
    invalid[column] = value  # 键已存在，赋值不改变 dict 顺序，values() 与表头列序仍一一对应

    content = "\n".join(
        [
            ",".join(invalid),
            ",".join(invalid.values()),
            "云器科技,后端开发,南京,,内推,,",  # 合法行：应正常入库
        ]
    ).encode("utf-8")

    result = application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert result.success_count == 1
    assert [(item.row, item.reason) for item in result.errors] == [(2, reason)]
    assert db_session.scalar(select(Application).where(Application.company == "云器科技")) is not None


def test_import_uppercase_status(db_session: Session, account_id: int):
    """TC-03：导入的 status 大小写不敏感，统一按大写入库。"""
    content = "company,position,status\n招银网络,后端开发,written\n".encode("utf-8")

    result = application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert result.success_count == 1
    assert result.errors == []
    row = db_session.scalar(select(Application).where(Application.company == "招银网络"))
    assert row.status == "WRITTEN"


def test_import_blank_optional_fields_to_none(db_session: Session, account_id: int):
    """TC-03：可选列留空入库为 NULL，而非空字符串。"""
    content = "company,position,city,channel,remark\n哈啰出行,后端开发,,,\n".encode("utf-8")

    result = application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert result.success_count == 1
    row = db_session.scalar(select(Application).where(Application.company == "哈啰出行"))
    assert (row.city, row.channel, row.remark) == (None, None, None)


# ---------- TC-95~97：导入模板中文表头（步骤 4 遗留，台账 #37） ----------

TEMPLATE_HEADER = "公司,岗位,城市,投递日期,渠道,状态,备注"


def test_import_chinese_header_succeeds(db_session: Session, account_id: int):
    """TC-95：模板口径的中文表头可正常导入，字段逐一映射正确。"""
    content = (
        f"{TEMPLATE_HEADER}\n"
        "满帮,Java 开发,南京,2026-09-01,官网,已投递,内推\n"
    ).encode("utf-8")

    result = application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert (result.success_count, result.errors) == (1, [])
    row = db_session.scalar(select(Application).where(Application.company == "满帮"))
    assert (row.company, row.position, row.city, row.channel) == ("满帮", "Java 开发", "南京", "官网")
    assert row.status == ApplicationStatus.APPLIED
    assert row.applied_at.date() == date(2026, 9, 1)  # 字段为 DateTime，比日期部分
    assert row.remark == "内推"


def test_import_skips_template_sample_row(db_session: Session, account_id: int):
    """TC-96：上传**未改动**的模板 → 示例行按首列标记跳过，0 条成功且 0 条错误。

    示例行若被当记录导入，会凭空多出一条「【示例】某某科技」——这正是标记要防的。
    """
    content = application_service.build_import_template()

    result = application_service.import_from_file(db_session, account_id, "template.xlsx", content)

    assert (result.success_count, result.errors) == (0, [])
    assert db_session.query(Application).count() == 0


def test_import_chinese_status_label(db_session: Session, account_id: int):
    """TC-97：状态列填中文标签与填英文枚举等效（模板里给的就是中文标签）。"""
    content = (
        f"{TEMPLATE_HEADER}\n"
        "亚信安全,后端开发,南京,2026-09-02,官网,面试中,\n"
        "浩鲸科技,Java 开发,南京,2026-09-03,官网,INTERVIEW,\n"
    ).encode("utf-8")

    result = application_service.import_from_file(db_session, account_id, "import.csv", content)

    assert (result.success_count, result.errors) == (2, [])
    rows = list(db_session.scalars(select(Application).order_by(Application.id)))
    assert [row.status for row in rows] == [ApplicationStatus.INTERVIEW] * 2


def test_download_template_has_chinese_header_and_sample(client: TestClient):
    """TC-98：下载模板 → xlsx 内容，表头为中文且含一行示例（首列带标记）。"""
    resp = client.get(f"{API}/template")

    assert resp.status_code == 200
    sheet = load_workbook(BytesIO(resp.content)).active
    header = [cell.value for cell in sheet[1]]
    assert header[:4] == ["公司", "岗位", "城市", "投递日期"]
    assert sheet.cell(row=2, column=1).value.startswith("【示例】")
