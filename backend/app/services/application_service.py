"""投递业务服务：CRUD、状态机校验、批量导入解析、趋势统计（FR-002~FR-005）。

分层规则（系统设计 3.1）：本层负责业务逻辑与事务边界，不返回 ORM 对象给路由层，出口一律为 schemas 层 DTO。
"""

import csv
from collections.abc import Sequence
from datetime import date, datetime, timedelta
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select, update
from sqlalchemy.orm import Session

from app.exceptions import BizException, ErrorCode
from app.models import Application, InterviewSession, JdAnalysisReport
from app.models.enums import ApplicationStatus, CloseReason
from app.schemas.application import (
    ApplicationCreate,
    ApplicationDTO,
    ApplicationListItem,
    ApplicationStatusUpdate,
    ImportErrorItem,
    ImportResult,
    TrendData,
    TrendItem,
)
from app.schemas.common import PageData
from app.utils.datetime_utils import DATE_FORMAT, date_range, to_datetime

# 合法状态流转（SRS FR-004，v1.8 起放开跳级）：链上顺序 已投递→待笔试→面试中→已获 offer，可前进到任意更靠后的
# 状态（跳过环节合法，如不设笔试的公司可从已投递直接进面试中），亦可终止；除已结束外任意状态均可流转至已结束
# （含 OFFER，覆盖拒 offer / 被撤回 / 谈崩），不可回退
LEGAL_TRANSITIONS: dict[ApplicationStatus, frozenset[ApplicationStatus]] = {
    ApplicationStatus.APPLIED: frozenset(
        {ApplicationStatus.WRITTEN, ApplicationStatus.INTERVIEW, ApplicationStatus.OFFER, ApplicationStatus.CLOSED}
    ),
    ApplicationStatus.WRITTEN: frozenset(
        {ApplicationStatus.INTERVIEW, ApplicationStatus.OFFER, ApplicationStatus.CLOSED}
    ),
    ApplicationStatus.INTERVIEW: frozenset({ApplicationStatus.OFFER, ApplicationStatus.CLOSED}),
    ApplicationStatus.OFFER: frozenset({ApplicationStatus.CLOSED}),
    ApplicationStatus.CLOSED: frozenset(),
}

# 状态中文名（数据库设计 §5），仅用于错误提示文案
STATUS_LABELS: dict[ApplicationStatus, str] = {
    ApplicationStatus.APPLIED: "已投递",
    ApplicationStatus.WRITTEN: "待笔试",
    ApplicationStatus.INTERVIEW: "面试中",
    ApplicationStatus.OFFER: "已获 offer",
    ApplicationStatus.CLOSED: "已结束",
}

# 结束原因中文名（数据库设计 §5），仅用于错误提示文案
CLOSE_REASON_LABELS: dict[CloseReason, str] = {
    CloseReason.FAILED: "未通过",
    CloseReason.DECLINED: "主动放弃",
    CloseReason.EXPIRED: "无消息",
}

# 导入模板表头（接口文档 3.3）与必填列
IMPORT_HEADERS = ("company", "position", "city", "applied_at", "channel", "status", "remark")
REQUIRED_HEADERS = ("company", "position")
IMPORTABLE_SUFFIXES = (".xlsx", ".csv")
MAX_IMPORT_BYTES = 2 * 1024 * 1024  # 上传文件上限 2MB（接口文档 3.3）


# ---------- 查询 ----------


def list_applications(
    db: Session,
    *,
    status: ApplicationStatus | None = None,
    close_reason: CloseReason | None = None,
    city: str | None = None,
    company: str | None = None,
    page: int = 1,
    page_size: int = 10,
) -> PageData[ApplicationListItem]:
    """投递列表：状态/结束原因/城市/公司关键字筛选 + 分页，按投递日期倒序（最近投递在前）。"""
    conditions = []
    if status is not None:
        conditions.append(Application.status == status)
    if close_reason is not None:
        conditions.append(Application.close_reason == close_reason)
    if city:
        conditions.append(Application.city == city)
    if company:
        conditions.append(Application.company.like(f"%{company}%"))

    total = db.scalar(select(func.count()).select_from(Application).where(*conditions)) or 0
    rows = db.scalars(
        select(Application)
        .where(*conditions)
        .order_by(Application.applied_at.desc(), Application.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return PageData[ApplicationListItem](total=total, items=[_to_list_item(row) for row in rows])


def get_application(db: Session, application_id: int) -> ApplicationDTO:
    """投递详情；不存在返回 10002。"""
    return _to_dto(_get_or_raise(db, application_id))


def trend(db: Session, days: int) -> TrendData:
    """投递趋势：按投递日期 GROUP BY 计数，补齐无投递日期（count=0），按日期升序（接口文档 3.3）。"""
    end = date.today()
    start = end - timedelta(days=days - 1)
    rows = db.execute(
        select(func.date(Application.applied_at), func.count())
        .where(Application.applied_at >= to_datetime(start))
        .group_by(func.date(Application.applied_at))
    ).all()
    counts = {str(day): count for day, count in rows}
    return TrendData(
        items=[
            TrendItem(date=day, count=counts.get(day.strftime(DATE_FORMAT), 0))
            for day in date_range(end, days)
        ]
    )


# ---------- 写入 ----------


def create_application(db: Session, payload: ApplicationCreate) -> ApplicationDTO:
    """新增投递：status 缺省 APPLIED、applied_at 缺省当天（接口文档 3.3）。"""
    now = datetime.now()
    status = payload.status or ApplicationStatus.APPLIED
    app = Application(
        company=payload.company,
        position=payload.position,
        city=payload.city,
        expected_salary=payload.expected_salary,
        applied_at=to_datetime(payload.applied_at or date.today()),
        channel=payload.channel,
        status=status,
        # 结束原因仅在记录处于 CLOSED 时有值（数据库设计 3.1）
        close_reason=payload.close_reason if status == ApplicationStatus.CLOSED else None,
        next_event_at=payload.next_event_at,
        remark=payload.remark,
        created_at=now,
        updated_at=now,
    )
    db.add(app)
    db.commit()
    db.refresh(app)
    return _to_dto(app)


def update_application(db: Session, application_id: int, payload: ApplicationCreate) -> ApplicationDTO:
    """编辑投递：全量更新，未传字段按空处理；status 不在本接口变更（走 PATCH，接口文档 3.3）。

    结束原因仅在记录当前处于 CLOSED 时可改（用于更正选错的原因），其余状态下传该字段返回 10001。
    """
    app = _get_or_raise(db, application_id)
    is_closed = ApplicationStatus(app.status) == ApplicationStatus.CLOSED
    if payload.close_reason is not None and not is_closed:
        raise BizException(ErrorCode.PARAM_INVALID, "仅已结束的投递可修改结束原因")
    app.company = payload.company
    app.position = payload.position
    app.city = payload.city
    app.expected_salary = payload.expected_salary
    app.applied_at = to_datetime(payload.applied_at or date.today())
    app.channel = payload.channel
    if is_closed:
        app.close_reason = payload.close_reason
    app.next_event_at = payload.next_event_at
    app.remark = payload.remark
    app.updated_at = datetime.now()
    db.commit()
    db.refresh(app)
    return _to_dto(app)


def delete_application(db: Session, application_id: int) -> None:
    """删除投递：物理删除（数据库设计 §1 不引入软删除）。

    关联的 JD 分析报告、模拟面试会话保留数据但解除关联（application_id 置空，两表该列本就可空），
    否则外键约束会拦下删除导致 500。提醒表 ref_id 无外键约束，不作为拦截源。
    """
    app = _get_or_raise(db, application_id)
    for model in (JdAnalysisReport, InterviewSession):
        db.execute(update(model).where(model.application_id == application_id).values(application_id=None))
    db.delete(app)
    db.commit()


def change_status(db: Session, application_id: int, payload: ApplicationStatusUpdate) -> ApplicationDTO:
    """状态流转：按 SRS FR-004 状态机校验，非法流转返回 10001；转 CLOSED 必须带结束原因。"""
    app = _get_or_raise(db, application_id)
    current = ApplicationStatus(app.status)
    target = payload.status
    if target not in LEGAL_TRANSITIONS[current]:
        raise BizException(
            ErrorCode.PARAM_INVALID,
            f"非法状态流转：{STATUS_LABELS[current]} → {STATUS_LABELS[target]}",
        )
    if target == ApplicationStatus.CLOSED and payload.close_reason is None:
        raise BizException(ErrorCode.PARAM_INVALID, "流转至已结束必须指明结束原因（未通过/主动放弃/无消息）")

    now = datetime.now()
    app.status = target
    # 结束原因仅在 CLOSED 态有值：转其他状态时清空（数据库设计 3.1）
    app.close_reason = payload.close_reason if target == ApplicationStatus.CLOSED else None
    # event_at 仅在转为 WRITTEN/INTERVIEW 时更新下次考试/面试时间（接口文档 3.3）
    if payload.event_at is not None and target in (ApplicationStatus.WRITTEN, ApplicationStatus.INTERVIEW):
        app.next_event_at = payload.event_at
    # 流转备注追加进投递备注（库中无流转记录表，追加保留既有内容不丢）
    if payload.remark:
        note = f"[{now.strftime('%Y-%m-%d %H:%M')} 流转至 {STATUS_LABELS[target]}] {payload.remark}"
        app.remark = f"{app.remark}\n{note}" if app.remark else note
    app.updated_at = now
    db.commit()
    db.refresh(app)
    return _to_dto(app)


# ---------- 批量导入 ----------


def import_from_file(db: Session, filename: str, content: bytes) -> ImportResult:
    """批量导入 xlsx/csv：行级校验互不影响，非法行进 errors；全部行非法时抛 20002（带 errors 清单）。"""
    _check_import_file(filename, content)
    header, rows = _parse_import_file(filename, content)
    _check_header(header)

    errors: list[ImportErrorItem] = []
    success_count = 0
    for row_no, raw in rows:
        if not any(value.strip() for value in raw.values()):
            continue  # 整行空白：跳过，不计入成功也不计为错误
        reason = _validate_row(raw)
        if reason is not None:
            errors.append(ImportErrorItem(row=row_no, reason=reason))
            continue
        db.add(_build_import_record(raw))
        success_count += 1

    if success_count == 0 and errors:
        db.rollback()
        raise BizException(
            ErrorCode.IMPORT_ROW_INVALID,
            f"共 {len(errors)} 行数据全部非法，未导入任何记录",
            data=ImportResult(success_count=0, errors=errors).model_dump(mode="json"),
        )

    db.commit()
    return ImportResult(success_count=success_count, errors=errors)


def build_import_template() -> bytes:
    """生成导入模板 xlsx：仅表头行（不预置示例行，避免示例数据被误当记录导入）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "投递清单"
    sheet.append(list(IMPORT_HEADERS))
    for index, name in enumerate(IMPORT_HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(name) + 6)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _check_import_file(filename: str, content: bytes) -> None:
    """文件级校验：扩展名、空文件、大小上限，不符合返回 20001。"""
    if Path(filename).suffix.lower() not in IMPORTABLE_SUFFIXES:
        raise BizException(ErrorCode.IMPORT_FILE_INVALID, "仅支持 .xlsx 或 .csv 文件")
    if not content:
        raise BizException(ErrorCode.IMPORT_FILE_INVALID, "文件内容为空")
    if len(content) > MAX_IMPORT_BYTES:
        raise BizException(ErrorCode.IMPORT_FILE_INVALID, "文件大小超过 2MB")


def _parse_import_file(filename: str, content: bytes) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    """按扩展名解析文件，返回 (表头, [(行号, {列名: 文本})])；行号从 2 起（第 1 行为表头）。"""
    if Path(filename).suffix.lower() == ".xlsx":
        return _parse_xlsx(content)
    return _parse_csv(content)


def _parse_xlsx(content: bytes) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    """解析 xlsx：取首个工作表的首行为表头。"""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError) as exc:
        raise BizException(ErrorCode.IMPORT_FILE_INVALID, "xlsx 文件无法解析，请下载模板重新填写") from exc
    sheet = workbook.active
    sheets = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not sheets:
        raise BizException(ErrorCode.IMPORT_FILE_INVALID, "文件为空，缺少表头行")
    header = [_cell_text(cell) for cell in sheets[0]]
    rows = [(row_no, _row_dict(header, values)) for row_no, values in enumerate(sheets[1:], start=2)]
    return header, rows


def _parse_csv(content: bytes) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    """解析 csv：首行为表头，编码优先 UTF-8（含 BOM），回退 GBK。"""
    reader = csv.reader(StringIO(_decode_csv(content)))
    sheets = list(reader)
    if not sheets:
        raise BizException(ErrorCode.IMPORT_FILE_INVALID, "文件为空，缺少表头行")
    header = [cell.strip() for cell in sheets[0]]
    rows = [(row_no, _row_dict(header, values)) for row_no, values in enumerate(sheets[1:], start=2)]
    return header, rows


def _decode_csv(content: bytes) -> str:
    """CSV 解码：Excel 另存 CSV 常见 GBK，故 UTF-8 失败后回退。"""
    for encoding in ("utf-8-sig", "gbk"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise BizException(ErrorCode.IMPORT_FILE_INVALID, "CSV 文件编码无法识别，请另存为 UTF-8 后重试")


def _row_dict(header: Sequence[str], values: Sequence) -> dict[str, str]:
    """按表头把一行单元格映射为 {列名: 文本}；表头之外的多余列忽略，缺列按空串。"""
    return {
        name: _cell_text(values[index]) if index < len(values) else ""
        for index, name in enumerate(header)
        if name
    }


def _cell_text(value) -> str:
    """单元格值 → 文本：Excel 日期单元格会读成 datetime/date，统一转为 YYYY-MM-DD。"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime(DATE_FORMAT)
    if isinstance(value, date):
        return value.strftime(DATE_FORMAT)
    return str(value).strip()


def _check_header(header: Sequence[str]) -> None:
    """表头校验：缺少必填列返回 20001（TC-04）。"""
    missing = [name for name in REQUIRED_HEADERS if name not in header]
    if missing:
        raise BizException(
            ErrorCode.IMPORT_FILE_INVALID,
            f"表头缺少必填列：{'、'.join(missing)}，请使用模板（{','.join(IMPORT_HEADERS)}）",
        )


def _validate_row(raw: dict[str, str]) -> str | None:
    """行级校验：返回 None 表示合法；否则返回非法原因（该行不入库，不影响其余行）。"""
    company = raw.get("company", "").strip()
    if not company:
        return "公司名称不能为空"
    if len(company) > 100:
        return "公司名称超过 100 字"

    position = raw.get("position", "").strip()
    if not position:
        return "岗位名称不能为空"
    if len(position) > 100:
        return "岗位名称超过 100 字"

    for field, label in (("city", "城市"), ("channel", "投递渠道")):
        if len(raw.get(field, "").strip()) > 50:
            return f"{label}超过 50 字"

    applied_at = raw.get("applied_at", "").strip()
    if applied_at:
        try:
            _parse_import_date(applied_at)
        except ValueError:
            return "投递日期格式应为 YYYY-MM-DD"

    status = raw.get("status", "").strip()
    if status:
        try:
            ApplicationStatus(status.upper())
        except ValueError:
            return f"状态值非法（可选：{'/'.join(item.value for item in ApplicationStatus)}）"
    return None


def _build_import_record(raw: dict[str, str]) -> Application:
    """把已通过校验的一行转成 ORM 记录：applied_at 缺省当天、status 缺省 APPLIED。"""
    now = datetime.now()
    applied_at = raw.get("applied_at", "").strip()
    status = raw.get("status", "").strip()
    return Application(
        company=raw["company"].strip(),
        position=raw["position"].strip(),
        city=raw.get("city", "").strip() or None,
        applied_at=to_datetime(_parse_import_date(applied_at)) if applied_at else to_datetime(date.today()),
        channel=raw.get("channel", "").strip() or None,
        status=ApplicationStatus(status.upper()) if status else ApplicationStatus.APPLIED,
        remark=raw.get("remark", "").strip() or None,
        created_at=now,
        updated_at=now,
    )


def _parse_import_date(text: str) -> date:
    """导入行的投递日期解析：接受 YYYY-MM-DD / YYYY/MM/DD，带时间部分则取日期。"""
    head = text.split(" ")[0]
    for fmt in (DATE_FORMAT, "%Y/%m/%d"):
        try:
            return datetime.strptime(head, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


# ---------- 内部转换 ----------


def _get_or_raise(db: Session, application_id: int) -> Application:
    """取投递记录，不存在返回 10002。"""
    app = db.get(Application, application_id)
    if app is None:
        raise BizException(ErrorCode.NOT_FOUND, "投递记录不存在")
    return app


def _to_dto(app: Application) -> ApplicationDTO:
    """ORM → 详情 DTO（applied_at 对外口径为日期，接口文档 3.3）。"""
    return ApplicationDTO(
        id=app.id,
        company=app.company,
        position=app.position,
        city=app.city,
        expected_salary=app.expected_salary,
        applied_at=app.applied_at.date(),
        channel=app.channel,
        status=ApplicationStatus(app.status),
        close_reason=CloseReason(app.close_reason) if app.close_reason else None,
        next_event_at=app.next_event_at,
        remark=app.remark,
        created_at=app.created_at,
        updated_at=app.updated_at,
    )


def _to_list_item(app: Application) -> ApplicationListItem:
    """ORM → 列表项 DTO（不含 remark）。"""
    return ApplicationListItem(**_to_dto(app).model_dump())
