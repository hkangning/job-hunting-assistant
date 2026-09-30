"""简历文件解析（FR-012）：PDF / Word 提取纯文本 + AI 结构化抽取画像字段。

**纯解析不落库**——解析结果回给前端填入表单，用户核对修改后走 `PUT /profile` 保存。

AI 抽取失败（未配 Key / 调用失败 / 输出异常）**不报错**：仍返回简历全文、`extracted` 记 null
（「部分成功也是成功」——报错会让用户连全文都拿不到）。文件本身解析失败或提取不到文字才报 10001。
"""

import io
import logging

import docx
import pdfplumber
from sqlalchemy.orm import Session

from app.clients.llm_client import LLMClient, resolve_config
from app.exceptions import BizException, ErrorCode
from app.prompts import build_resume_extract_messages
from app.schemas.system import ResumeExtracted, ResumeParseData

logger = logging.getLogger(__name__)

MAX_RESUME_BYTES = 10 * 1024 * 1024  # 上传体积上限 10MB（接口文档 3.12）
RESUME_EXTS = (".pdf", ".docx")

# 抽取字段的截断上限（与 ProfileUpdateRequest 的各字段 max_length 一致）：
# 超长会因前端回填保存被 422 拦下，故在解析侧先收口
_FIELD_LIMITS = {"name": 50, "school": 100, "major": 100, "degree": 20, "gpa": 20, "english_level": 50}
_SKILL_LIMIT = 30  # 单个技能词长度上限
_SKILLS_MAX = 20  # 技能项数上限（prompt 要求 3~10 项，此处放宽做兜底）


def parse_resume(
    db: Session, *, user_id: int, filename: str, content: bytes, client: LLMClient
) -> ResumeParseData:
    """解析上传的简历文件（POST /profile/resume）。

    扩展名 / 体积 / 可解析性校验失败抛 10001；AI 抽取失败不抛错、`extracted` 记 null。
    """
    ext = _suffix(filename)
    if ext not in RESUME_EXTS:
        raise BizException(ErrorCode.PARAM_INVALID, "仅支持 PDF 或 Word（.docx）格式的简历")
    if not content:
        raise BizException(ErrorCode.PARAM_INVALID, "文件内容为空")
    if len(content) > MAX_RESUME_BYTES:
        raise BizException(ErrorCode.PARAM_INVALID, "文件不能超过 10MB")

    text = _extract_pdf(content) if ext == ".pdf" else _extract_docx(content)
    text = _tidy(text)
    if not text:
        raise BizException(ErrorCode.PARAM_INVALID, "未能从文件中提取到文字，若为扫描件 PDF 请改用文字版")
    return ResumeParseData(
        resume_text=text,
        extracted=_extract_fields(db, user_id=user_id, text=text, client=client),
    )


def _extract_pdf(content: bytes) -> str:
    """PDF 文字层逐页提取；文件损坏 / 加密等解析失败统一转 10001。"""
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as exc:
        raise BizException(ErrorCode.PARAM_INVALID, "PDF 解析失败，请确认文件未损坏、未加密") from exc


def _extract_docx(content: bytes) -> str:
    """Word 提取：段落按行 + 表格逐行（简历常用表格排版，漏掉表格会丢大段内容）。"""
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise BizException(ErrorCode.PARAM_INVALID, "Word 文档解析失败，请确认文件格式正确") from exc

    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:  # 合并单元格会重复出现同一文本，按行去重
                text = cell.text.strip().replace("\n", " ")
                if text and text not in cells:
                    cells.append(text)
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _tidy(text: str) -> str:
    """压掉连续空行与行尾空白（PDF 提取常带大量空行，直接入库会让画像文本难看）。"""
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    tidied: list[str] = []
    for line in lines:
        if not line and (not tidied or not tidied[-1]):
            continue
        tidied.append(line)
    return "\n".join(tidied).strip()


def _extract_fields(
    db: Session, *, user_id: int, text: str, client: LLMClient
) -> ResumeExtracted | None:
    """AI 结构化抽取：任何失败（未配 Key / 调用失败 / 输出异常）都返回 null，不阻断全文返回。"""
    try:
        config = resolve_config(db, user_id)
        result = client.chat_json(config, build_resume_extract_messages(text))
    except Exception:
        # 含未配 Key（10012）与模型输出异常（10011）：降级为「只给全文」，前端据此不填表单
        logger.warning("简历字段抽取失败（账号 %s），降级为仅返回全文", user_id, exc_info=True)
        return None
    return _clean(result)


def _clean(result: dict) -> ResumeExtracted | None:
    """字段清洗：按画像字段长度上限截断、skills 归一为去重字符串数组；全空时返回 null。"""
    values: dict = {}
    for field, limit in _FIELD_LIMITS.items():
        values[field] = _as_text(result.get(field))[:limit] or None
    skills = result.get("skills")
    cleaned: list[str] = []
    if isinstance(skills, list):
        for item in skills:
            name = _as_text(item)[:_SKILL_LIMIT]
            if name and name not in cleaned:
                cleaned.append(name)
    values["skills"] = cleaned[:_SKILLS_MAX] or None
    if not any(values.values()):
        return None
    return ResumeExtracted(**values)


def _as_text(value: object) -> str:
    """抽取值归一为字符串：字符串原样、数字转文本（模型偶尔给数字），其余视为无效。"""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(value)
    return ""


def _suffix(filename: str) -> str:
    """取小写扩展名（含点）；无扩展名返回空串。"""
    name = (filename or "").strip().lower()
    dot = name.rfind(".")
    return name[dot:] if dot >= 0 else ""
